from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class DocumentService:
    """Build compact, professional Word documents entirely in memory."""

    BLUE = "1F5FBF"
    LIGHT_BLUE = "EAF2FD"
    TEXT = RGBColor(24, 44, 76)
    MUTED = RGBColor(92, 111, 139)
    FONT = "Noto Sans CJK SC"

    def generate(
        self,
        *,
        title: str,
        subtitle: str | None,
        summary: str | None,
        metadata: list[dict[str, str]],
        sections: list[dict[str, Any]],
    ) -> bytes:
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

        self._configure_styles(document)
        self._add_top_rule(document)

        title_paragraph = document.add_paragraph(style="Title")
        title_paragraph.add_run(title.strip())
        if subtitle:
            subtitle_paragraph = document.add_paragraph()
            subtitle_paragraph.style = document.styles["Subtitle"]
            subtitle_paragraph.add_run(subtitle.strip())

        if metadata:
            self._add_metadata_table(document, metadata)

        if summary:
            document.add_heading("概览", level=1)
            self._add_body_paragraph(document, summary)

        for item in sections:
            heading = str(item.get("heading") or "").strip()
            if heading:
                document.add_heading(heading, level=1)
            for paragraph in item.get("paragraphs") or []:
                self._add_body_paragraph(document, str(paragraph))
            for bullet in item.get("bullets") or []:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(str(bullet).strip())
            numbered = item.get("numbered_items") or []
            for entry in numbered:
                paragraph = document.add_paragraph(style="List Number")
                paragraph.add_run(str(entry).strip())
            rows = item.get("table") or []
            if rows:
                self._add_data_table(document, rows)

        self._add_footer(document)
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def _configure_styles(self, document: Document) -> None:
        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = self.FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = self.TEXT
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        self._set_east_asia_font(normal, self.FONT)

        title = styles["Title"]
        title.font.name = self.FONT
        title.font.size = Pt(25)
        title.font.bold = True
        title.font.color.rgb = self.TEXT
        title.paragraph_format.space_before = Pt(12)
        title.paragraph_format.space_after = Pt(5)
        self._set_east_asia_font(title, self.FONT)

        subtitle = styles["Subtitle"]
        subtitle.font.name = self.FONT
        subtitle.font.size = Pt(10.5)
        subtitle.font.color.rgb = self.MUTED
        subtitle.paragraph_format.space_after = Pt(14)
        self._set_east_asia_font(subtitle, self.FONT)

        heading = styles["Heading 1"]
        heading.font.name = self.FONT
        heading.font.size = Pt(15)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(6)
        heading.paragraph_format.keep_with_next = True
        self._set_east_asia_font(heading, self.FONT)

        for style_name in ("List Bullet", "List Number"):
            style = styles[style_name]
            style.font.name = self.FONT
            style.font.size = Pt(10.5)
            style.font.color.rgb = self.TEXT
            style.paragraph_format.space_after = Pt(4)
            self._set_east_asia_font(style, self.FONT)

    def _add_top_rule(self, document: Document) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        properties = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "18")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), self.BLUE)
        borders.append(bottom)
        properties.append(borders)

    def _add_metadata_table(
        self, document: Document, metadata: list[dict[str, str]]
    ) -> None:
        table = document.add_table(rows=0, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.25)
        table.columns[1].width = Inches(5.55)
        for item in metadata:
            row = table.add_row()
            row.cells[0].width = Inches(1.25)
            row.cells[1].width = Inches(5.55)
            label_cell, value_cell = row.cells
            label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self._shade_cell(label_cell, self.LIGHT_BLUE)
            label_run = label_cell.paragraphs[0].add_run(item["label"].strip())
            label_run.bold = True
            value_cell.paragraphs[0].add_run(item["value"].strip())
        document.add_paragraph().paragraph_format.space_after = Pt(0)

    def _add_data_table(self, document: Document, rows: list[list[str]]) -> None:
        rows = [row for row in rows if row]
        if not rows:
            return
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        table = document.add_table(rows=len(normalized), cols=width)
        table.style = "Table Grid"
        for row_index, row in enumerate(normalized):
            for column_index, value in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                run = cell.paragraphs[0].add_run(str(value).strip())
                if row_index == 0:
                    self._shade_cell(cell, self.BLUE)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        document.add_paragraph().paragraph_format.space_after = Pt(0)

    @staticmethod
    def _add_body_paragraph(document: Document, text: str) -> None:
        cleaned = text.strip()
        if cleaned:
            document.add_paragraph(cleaned)

    def _add_footer(self, document: Document) -> None:
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("FlowAgent 自动生成")
        run.font.name = self.FONT
        run.font.size = Pt(8)
        run.font.color.rgb = self.MUTED

    @staticmethod
    def _set_east_asia_font(style, font_name: str) -> None:
        style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    @staticmethod
    def _shade_cell(cell, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shade = properties.find(qn("w:shd"))
        if shade is None:
            shade = OxmlElement("w:shd")
            properties.append(shade)
        shade.set(qn("w:fill"), fill)
